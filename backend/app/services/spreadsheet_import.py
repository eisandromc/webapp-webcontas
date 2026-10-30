from __future__ import annotations

import hashlib
import io
import re
from dataclasses import asdict, dataclass, field
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP

import openpyxl


MONTH_RE = re.compile(r"^(\d{4})-(\d{2})_")
FIXED_CATEGORIES = {"LUZ", "AGUA", "ÁGUA", "GAS", "INTERNET", "COND.", "CONDOMINIO", "CONDOMÍNIO"}


def cents(value) -> int:
    return int((Decimal(str(value)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def clean_text(value) -> str:
    return " ".join(str(value or "").replace("�", "Ê").split()).strip()


def source_date(value, period: str) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = clean_text(value)
    match = re.search(r"(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?", text)
    if not match:
        return None
    day, month = int(match.group(1)), int(match.group(2))
    year = int(period[:4])
    if match.group(3):
        year = int(match.group(3))
        if year < 100:
            year += 2000
    try:
        return datetime(year, month, day).date().isoformat()
    except ValueError:
        return None


@dataclass
class ParsedRecord:
    kind: str
    period: str
    source_key: str
    description: str
    amount_cents: int
    category_slug: str | None = None
    date: str | None = None
    adjustment_kind: str | None = None


@dataclass
class ParsedIssue:
    period: str | None
    sheet: str
    cell: str | None
    code: str
    message: str
    source_value: str | None = None


@dataclass
class SpreadsheetPreview:
    file_hash: str
    months: list[str] = field(default_factory=list)
    records: list[ParsedRecord] = field(default_factory=list)
    issues: list[ParsedIssue] = field(default_factory=list)

    def public(self) -> dict:
        per_month: dict[str, dict] = {}
        for record in self.records:
            item = per_month.setdefault(record.period, {"period": record.period, "entries": 0, "adjustments": 0})
            item["entries" if record.kind == "entry" else "adjustments"] += 1
        return {
            "file_hash": self.file_hash,
            "months": [per_month[period] for period in sorted(per_month)],
            "record_count": len(self.records),
            "issue_count": len(self.issues),
            "issues": [asdict(issue) for issue in self.issues],
        }


def parse_spreadsheet(data: bytes) -> SpreadsheetPreview:
    digest = hashlib.sha256(data).hexdigest()
    formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    values = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    preview = SpreadsheetPreview(file_hash=digest)

    for ws in values.worksheets:
        match = MONTH_RE.match(ws.title)
        if not match:
            preview.issues.append(ParsedIssue(None, ws.title, None, "unknown_sheet", "A aba não segue o padrão AAAA-MM_Nome."))
            continue
        period = f"{match.group(1)}-{match.group(2)}"
        preview.months.append(period)
        header = ws["A1"].value if ws["A1"].value is not None else ws["A2"].value
        if isinstance(header, datetime) and header.strftime("%Y-%m") != period:
            preview.issues.append(
                ParsedIssue(period, ws.title, "A1/A2", "header_month_mismatch", f"Cabeçalho {header:%Y-%m} difere da aba {period}.", str(header.date()))
            )

        for row in range(1, ws.max_row + 1):
            label = clean_text(ws.cell(row, 1).value)
            upper = label.upper()
            value = ws.cell(row, 2).value
            if upper in FIXED_CATEGORIES and isinstance(value, (int, float)):
                preview.records.append(
                    ParsedRecord("entry", period, f"xlsx:{digest}:{ws.title}:B{row}", label.title(), cents(value), "contas")
                )
            elif upper.startswith("PAGO") and isinstance(value, (int, float)):
                parsed = source_date(label, period)
                if parsed is None:
                    preview.issues.append(ParsedIssue(period, ws.title, f"A{row}:B{row}", "invalid_payment_date", "Pagamento legado precisa de data e pessoa na revisão.", f"{label}: {value}"))
                else:
                    preview.issues.append(ParsedIssue(period, ws.title, f"A{row}:B{row}", "payment_owner_required", "Escolha quem realizou este pagamento antes de aplicá-lo.", f"{label}: {value}"))
            elif "FALTOU" in upper and isinstance(value, (int, float)):
                preview.issues.append(ParsedIssue(period, ws.title, f"A{row}:B{row}", "legacy_balance_owner_required", "Saldo anterior legado precisa ser atribuído manualmente.", str(value)))

        for row in range(4, ws.max_row + 1):
            description = clean_text(ws.cell(row, 4).value)
            value = ws.cell(row, 5).value
            if not description or not isinstance(value, (int, float)):
                continue
            upper = description.upper()
            key = f"xlsx:{digest}:{ws.title}:E{row}"
            if "DESC VALE" in upper:
                preview.records.append(ParsedRecord("adjustment", period, key, description, cents(value), adjustment_kind="desconto"))
            elif upper == "VALE" or upper.startswith("VALE "):
                preview.records.append(ParsedRecord("adjustment", period, key, description, cents(value), adjustment_kind="vale"))
            elif not upper.startswith("TOTAL"):
                preview.records.append(ParsedRecord("entry", period, key, description, cents(value), "alimentacao", source_date(description, period)))

        current_extra_category = "extras"
        for row in range(4, ws.max_row + 1):
            description = clean_text(ws.cell(row, 6).value)
            value = ws.cell(row, 7).value
            upper = description.upper()
            if upper == "TRANSPORTE":
                current_extra_category = "transporte"
                continue
            if upper == "GASTOS" or not description or not isinstance(value, (int, float)):
                continue
            preview.records.append(
                ParsedRecord("entry", period, f"xlsx:{digest}:{ws.title}:G{row}", description, cents(value), current_extra_category, source_date(description, period))
            )

        for row in range(6, ws.max_row + 1):
            description = clean_text(ws.cell(row, 9).value)
            value = ws.cell(row, 10).value
            if description and isinstance(value, (int, float)):
                preview.records.append(
                    ParsedRecord("adjustment", period, f"xlsx:{digest}:{ws.title}:J{row}", description, cents(value), adjustment_kind="desconto")
                )

        formula_ws = formulas[ws.title]
        for row in formula_ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=") and "1048576" in cell.value:
                    preview.issues.append(ParsedIssue(period, ws.title, cell.coordinate, "unbounded_formula", "Fórmula antiga usa a coluna inteira; os valores serão recalculados.", cell.value))
        if ws.title == "2025-02_Fev" and formula_ws["C8"].value:
            preview.issues.append(
                ParsedIssue(period, ws.title, "C8", "known_division_anomaly", "O valor dividido da alimentação não corresponde à metade do total e não será preservado.", str(ws["C8"].value))
            )
    return preview

