from io import BytesIO

import openpyxl

from app.services.spreadsheet_import import parse_spreadsheet


def build_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2026-09_Set"
    ws["A2"] = "2026-10-01"
    ws["A4"], ws["B4"] = "LUZ", 200
    ws["D4"], ws["E4"] = "DESC VALE", 20
    ws["D5"], ws["E5"] = "VALE", 100
    ws["D6"], ws["E6"] = "Mercado 04/09", 50
    ws["F4"] = "GASTOS"
    ws["F5"], ws["G5"] = "Farmácia 05/09", 30
    ws["A10"], ws["B10"] = "Pago - 31/06", 100
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def test_spreadsheet_parser_uses_sheet_name_and_separates_adjustments():
    preview = parse_spreadsheet(build_workbook())
    assert preview.months == ["2026-09"]
    assert len([item for item in preview.records if item.kind == "entry"]) == 3
    assert {item.adjustment_kind for item in preview.records if item.kind == "adjustment"} == {"vale", "desconto"}
    assert any(issue.code == "invalid_payment_date" for issue in preview.issues)

